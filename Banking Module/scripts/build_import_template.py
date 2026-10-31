"""Builds the only Excel the CRM accepts for lead import (needs openpyxl, in the CRM .venv).
Keep COLUMNS, MARKER and VERSION in step with logic/leadTemplate.js; a test checks they match.
    .venv\\Scripts\\python.exe "Banking Module\\scripts\\build_import_template.py"
"""
import shutil
from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "Banking Module" / "templates" / "Lead_Import_Template.xlsx"
MASTER = ROOT / "Resources" / "Banking"
MARKER, VERSION, ROWS = "AB_CAPITAL_LEAD_TEMPLATE", 1, 2000

BANKS = ["ADCB", "ADIB", "Dubai Islamic Bank", "Emirates Islamic Bank", "Emirates NBD", "FAB", "M Bank", "Mashreq", "National Bank of Fujairah", "RAKBANK", "Ruya Bank", "Sharjah Islamic Bank", "UBL Bank", "Wio Bank"]
SOURCES = ["Cold Call", "WhatsApp", "Meta Ads", "Google Ads", "Instagram", "Direct", "Campaign", "Third Party", "Employee", "Referral", "Other"]
STATUSES = ["Document Collection", "Documents Under Review", "Documents Completed", "Application Preparation", "Ready for Submission", "Submitted to Bank", "Bank Review", "Additional Documents Requested", "Resubmission Required", "Appointment Scheduled", "Client Visited Bank", "Approved", "Account Opened", "Rejected", "Returned", "Cancelled", "On Hold"]
METHODS = ["Bank Transfer", "Cash", "Cash Deposit", "Cheque", "Card", "Other"]

# header, required, kind, list name, help
COLUMNS = [
    ("Company Name", True, "text", None, "Legal name of the company or client."),
    ("Contact Person", False, "text", None, None),
    ("Mobile", False, "text", None, None),
    ("WhatsApp", False, "text", None, None),
    ("Email", False, "text", None, None),
    ("Nationality", False, "text", None, None),
    ("License Number", False, "text", None, None),
    ("Bank", True, "list", "banks", "Choose from the list."),
    ("Account Type", True, "list", "kinds", "Business or Personal."),
    ("Lead Date", True, "date", None, "A real Excel date, not typed text."),
    ("Lead Source", True, "list", "sources", "Choose from the list."),
    ("Referred By", False, "text", None, None),
    ("Application Status", True, "list", "statuses", "Choose from the list."),
    ("Fee AED", False, "number", None, "Numbers only. Total agreed fee."),
    ("Home Ejari", False, "list", "yesno", "Yes only if the client took this service."),
    ("Relationship Manager", False, "text", None, None),
    ("Payment 1 Amount", False, "number", None, "Numbers only. Fill amount and date together."),
    ("Payment 1 Date", False, "date", None, None),
    ("Payment 1 Method", False, "list", "methods", None),
    ("Payment 2 Amount", False, "number", None, None),
    ("Payment 2 Date", False, "date", None, None),
    ("Payment 2 Method", False, "list", "methods", None),
    ("Payment 3 Amount", False, "number", None, None),
    ("Payment 3 Date", False, "date", None, None),
    ("Payment 3 Method", False, "list", "methods", None),
    ("Remarks", False, "text", None, None),
]
LISTS = {"banks": BANKS, "kinds": ["Business", "Personal"], "sources": SOURCES, "statuses": STATUSES, "yesno": ["Yes", "No"], "methods": METHODS}

wb = Workbook()
leads = wb.active
leads.title = "Leads"
navy, grey = PatternFill("solid", fgColor="13363D"), PatternFill("solid", fgColor="5F7077")

lists = wb.create_sheet("Lists")
refs = {}
for col, (name, values) in enumerate(LISTS.items(), start=1):
    letter = lists.cell(1, col, name).column_letter
    for row, value in enumerate(values, start=2):
        lists.cell(row, col, value)
    refs[name] = f"Lists!${letter}$2:${letter}${len(values) + 1}"
lists.sheet_state = "hidden"
lists.protection.sheet = True

for col, (header, required, kind, list_name, help_text) in enumerate(COLUMNS, start=1):
    cell = leads.cell(1, col, header)
    cell.font, cell.fill = Font(bold=True, color="FFFFFF"), navy if required else grey
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    if help_text:
        cell.comment = Comment(help_text, "AB Capital")
    letter = cell.column_letter
    leads.column_dimensions[letter].width = 26 if col in (1, 26) else 16
    span = f"{letter}2:{letter}{ROWS + 1}"
    dv = None
    if kind == "list":
        dv = DataValidation(type="list", formula1=f"={refs[list_name]}", allow_blank=not required, showErrorMessage=True, errorTitle=header, error="Choose a value from the list.")
    elif kind == "date":
        dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=not required, showErrorMessage=True, errorTitle=header, error="Enter a real date.")
    elif kind == "number":
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, errorTitle=header, error="Numbers only.")
    if dv:
        leads.add_data_validation(dv)
        dv.add(span)
    for row in range(2, ROWS + 2):
        c = leads.cell(row, col)
        c.protection = Protection(locked=False)
        if kind == "date":
            c.number_format = "dd-mmm-yyyy"
leads.row_dimensions[1].height = 32
leads.freeze_panes = "B2"
leads.protection.sheet = True  # headers stay fixed; no password, it guards against slips rather than intent
leads.protection.formatColumns = False

guide = wb.create_sheet("How to fill", 0)
for i, line in enumerate([
    "AB Capital lead import template",
    "1. Fill the Leads sheet, one lead per row. Dark blue columns are required.",
    "2. Choose list values from the dropdowns. Do not paste over them.",
    "3. Do not rename, add, move or delete columns or sheets. The CRM refuses a changed file.",
    "4. Enter payments only with both an amount and a date.",
    "5. Save, then upload in Business Banking > Excel Import.",
], start=1):
    guide.cell(i, 1, line).font = Font(bold=i == 1, size=14 if i == 1 else 11)
guide.column_dimensions["A"].width = 95

marker = wb.create_sheet("_template")
marker["A1"], marker["A2"] = MARKER, VERSION
marker.sheet_state = "veryHidden"
wb.active = 1
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
if MASTER.exists():
    shutil.copy(OUT, MASTER / OUT.name)
print("built", OUT.name)
