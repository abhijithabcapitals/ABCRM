"""Build SQL Data/BD_Field_Mapping.xlsx: every BD spec field mapped to its database column.

Sheets:
  Legend                 how to read the workbook
  Spec to Database       each field of docs/BD_BACKEND_SPEC.md -> table.column in migrations/bd/001_bd_schema.sql
  Client Master Prefill  BD data -> Client Module client-master.csv column at onboarding handover
  Reuse Decisions        existing CRM entities reused, linked or created

    "D:\\Abhijith Projects\\CRM\\.venv\\Scripts\\python.exe" build_bd_field_mapping.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
SPEC = HERE / "docs" / "BD_BACKEND_SPEC.md"
DDL = HERE / "migrations" / "bd" / "001_bd_schema.sql"
OUT = HERE / "SQL Data" / "BD_Field_Mapping.xlsx"

FONT = "Arial"
NAVY = "1F3A4D"
THIN = Side(style="thin", color="C9D3DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {"As spec": "E8F3EA", "Adjusted": "FFF4D6", "Added": "E3EDF7", "Missing": "F8D7DA"}

# Deliberate differences from the spec, explained where the reader sees them
NOTES = {
    ("bd_leads", "converted_client_id"): "VARCHAR(20) Client ID (CL000001), not UUID: the Client Master is the Client "
                                         "Module CSV. Format CHECK + API lookup via Client Module GET /api/clients/{id}.",
    ("bd_onboarding_handovers", "client_id"): "VARCHAR(20) Client ID (CL000001), not UUID - see bd_leads.converted_client_id.",
    ("bd_leads", "acquired_by_user_id"): "Trigger keeps the original attribution; reassignment changes assigned_to_user_id.",
    ("bd_leads", "case_opened_at"): "Trigger makes it immutable.",
    ("bd_leads", "lead_number"): "Default LD-YYYY-000001 from a sequence; immutable.",
    ("bd_leads", "assigned_team_id"): "FK to crm_teams (created: the CRM had no team entity).",
    ("bd_leads", "initial_discussion_summary"): "NOT NULL, at least 10 characters; API also rejects placeholder text.",
    ("bd_leads", "mobile_number"): "Duplicate checks use the generated column mobile_normalized (digits only).",
    ("bd_leads", "email"): "Duplicate checks use the generated column email_normalized (lower, trimmed).",
    ("bd_leads", "vat_trn"): "CHECK: 15 digits.",
    ("bd_leads", "status"): "CHECKs: Won needs won_at + closing_user_id; Lost needs lost_at + lost_reason_id; "
                            "On Hold needs hold_until.",
    ("bd_lead_services", "service_id"): "FK to bd_service_catalog (created: no service master existed).",
    ("bd_lead_services", "quotation_id"): "FK to bd_quotations.",
    ("bd_quotations", "quotation_number"): "Default QT-YYYY-000001; issued quotations are frozen (new version instead).",
    ("bd_onboarding_handovers", "handover_number"): "Default HO-YYYY-000001; one handover not cancelled per lead.",
    ("bd_discussions", "next_followup_at"): "CHECK: next_followup_at or no_followup_reason is required.",
    ("bd_discussions", "discussion_summary"): "NOT NULL, at least 10 characters; API rejects generic placeholders.",
    ("bd_followups", "status"): "CHECK: Completed needs completed_at, outcome and completion_summary.",
    ("bd_followups", "scheduled_at"): "Reschedules keep the original time in bd_followup_history.",
    ("bd_audit_logs", "action"): "Also allows Archive, Lost, Reopen, Export. Table is append-only (trigger).",
    ("bd_campaigns", "currency_code"): "CHECK: required when budget_amount or actual_spend is entered.",
}

# Columns added beyond the spec, with the reason
ADDED = {
    ("bd_leads", "mobile_normalized"): "Generated: digits of country code + mobile, for duplicate checks.",
    ("bd_leads", "email_normalized"): "Generated: lower(trim(email)), for duplicate checks.",
    ("bd_quotations", "version_no"): "Spec: preserve issued quotation versions.",
    ("bd_quotations", "supersedes_quotation_id"): "Links a new version to the one it replaces.",
    ("bd_quotations", "quotation_template_id"): "Freezone .docx template used (public.quotation_templates on the live CRM).",
    ("bd_onboarding_handovers", "idempotency_key"): "Mark Won retries return the same handover.",
    ("bd_lead_stage_history", "from_status"): "Status at the transition, so Won/Lost/Reopen show in the same history.",
    ("bd_lead_stage_history", "to_status"): "Status after the transition.",
}

CLIENT_PREFILL = [
    # (BD source, Client Module column, rule)
    ("bd_leads.company_name (Company) / prospect_name (Individual)", "Client Legal Name",
     "Marked (Provided – Verification Required) until checked against the trade licence."),
    ("bd_leads.entity_type", "Entity Type", "Only if it matches a Client Module option; otherwise left blank."),
    ("bd_leads.trade_licence_number", "Trade License Number", "Verification Required."),
    ("bd_leads.trade_licence_expiry_date", "License Expiry Date", "DD/MM/YYYY; Verification Required."),
    ("bd_leads.emirate", "Emirate", "Only if it matches a Client Module option."),
    ("bd_leads.website", "Website", "As captured."),
    ("bd_leads.email", "General Email", "As captured."),
    ("bd_leads.mobile_country_code + mobile_number", "General Phone", "As captured."),
    ("bd_leads.primary_contact_name, designation, email, mobile, whatsapp", "Contact Details",
     "One contact, Roles: Primary. Fills Primary Contact."),
    ("bd_leads.vat_trn", "TRN", "Verification Required. VAT Registered is NOT inferred from a TRN being present."),
    ("bd_leads.corporate_tax_status", "Tax Notes", "Text note 'BD: Corporate tax status ...'; CT fields not inferred."),
    ("bd_lead_services (Accepted) -> bd_service_catalog.service_name", "Services Provided",
     "'; '-separated. Catalog names equal Client Module options."),
    ("bd_lead_services.final_agreed_fee_amount where billing_frequency = Monthly (AED)", "Monthly Fee",
     "Sum of Monthly accepted AED fees, as 'AED 1,500.00'. Other currencies go to Management Notes, never added."),
    ("bd_lead_services.final_agreed_fee_amount where billing_frequency = Annual (AED)", "Annual Fee",
     "Same rule as Monthly Fee."),
    ("bd_lead_services.billing_frequency", "Billing Frequency",
     "Only when all accepted services share one frequency (One-time -> One-time, Annual -> Annually)."),
    ("bd_leads.initial_discussion_summary + latest discussion requirement_summary", "Business Notes",
     "Prefixed 'BD:'. Not copied into Business Model Summary (never written for the client)."),
    ("bd_leads.lead_number, source, closing user, won_at", "Management Notes",
     "'BD lead LD-2026-000001 · Source Meta Ads · Closed by <name> on DD/MM/YYYY'."),
    ("bd_onboarding_handovers.handover_notes, missing_items", "Important Instructions", "As written by the closer."),
    ("bd_leads.assigned_to_user_id (profiles.full_name)", "Account Manager", "Suggestion only; onboarding confirms."),
    ("(fixed)", "Onboarding Status", "'New'."),
    ("(Client Module)", "Client ID", "Issued by the Client Module on save; written back to bd_leads.converted_client_id."),
]

REUSE = [
    ("Users / employees", "public.profiles (live CRM, id = auth.users.id)", "Reuse",
     "All user FKs reference profiles(id). Local DB gets a same-shape stand-in."),
    ("Teams", "None", "Create", "crm_teams, crm_team_members (general CRM tables)."),
    ("Roles / permissions", "profiles.role admin/member", "Extend", "bd_user_roles for BD-specific roles."),
    ("Client Master", "Client Module data/client-master.csv (CL000001)", "Link",
     "converted_client_id / handover client_id hold the Client ID; clients created only via the Client Module."),
    ("Service catalog", "Client Module 'Services Provided' options + config.json serviceFolders", "Create",
     "bd_service_catalog seeded with the same 27 names and folders."),
    ("Quotations", "quotation_templates (freezone .docx templates); no quotation records", "Create",
     "bd_quotations + bd_quotation_items; quotation_template_id records the template used."),
    ("Notifications", "public.notifications", "Reuse", "Follow-up reminders written by the API."),
    ("Documents", "Document Module pipeline + DataBase folders", "Reuse",
     "attachment_reference stores a secure reference, never a public URL."),
    ("Currency / FX", "quotation_templates.usd_rate only", "Not created",
     "Amounts carry their currency; no cross-currency totals until an FX source is agreed."),
    ("Audit", "task_history, automation_runs (task-specific)", "Create", "bd_audit_logs, append-only."),
    ("Agent Module bizdev schema", "SQL Server design, not deployed", "Not followed",
     "Spec and live CRM are PostgreSQL (decision D1 in the migration plan)."),
]


def spec_fields() -> list[tuple[str, str, str, str]]:
    """(table, field, type, requirement) from the spec's 'Table N - bd_x' sections."""
    out, table = [], None
    for line in SPEC.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^## \d+\. Table \d+ — (bd_\w+)", line)
        if m:
            table = m.group(1)
            continue
        if line.startswith("## "):
            table = None
        if table and line.startswith("| ") and not line.startswith("| Field") and not line.startswith("|---"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if len(cells) >= 3:
                out.append((table, cells[0], cells[1], cells[2]))
    return out


def ddl_columns() -> dict[str, list[tuple[str, str]]]:
    """table -> [(column, definition)] from the CREATE TABLE blocks."""
    tables, cur = {}, None
    for line in DDL.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^create table if not exists public\.(\w+) \(", line)
        if m:
            cur = m.group(1)
            tables[cur] = []
            continue
        if cur and line.startswith(");"):
            cur = None
            continue
        if cur:
            m = re.match(r"^\s{2}(\w+)\s+(.+?),?\s*(--.*)?$", line)
            if m and m.group(1) not in ("constraint", "primary", "unique", "check"):
                tables[cur].append((m.group(1), m.group(2).rstrip(",")))
    return tables


def short_type(defn: str) -> str:
    m = re.match(r"(\w+(?:\s?\(\d+(?:,\d+)?\))?)", defn)
    t = m.group(1) if m else defn
    flags = []
    if "primary key" in defn:
        flags.append("PK")
    if "references" in defn:
        flags.append("FK " + re.search(r"references public\.(\w+)", defn).group(1))
    if re.search(r"\bunique\b", defn):
        flags.append("UNIQUE")
    if "not null" in defn:
        flags.append("NOT NULL")
    if "generated always" in defn:
        flags.append("GENERATED")
    if re.search(r"\bcheck\b", defn):
        flags.append("CHECK")
    if " default " in f" {defn} ":
        flags.append("DEFAULT")
    return t.upper() + (" · " + ", ".join(flags) if flags else "")


def header(ws, heads: list[str], widths: list[int]):
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(1, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def write_rows(ws, rows: list[list], wrap_cols: set[int], status_col: int | None = None):
    for r, row in enumerate(rows, start=2):
        for i, v in enumerate(row, start=1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FONT)
            c.alignment = Alignment(wrap_text=i in wrap_cols, vertical="top")
            c.border = BORDER
        if status_col:
            fill = STATUS_FILL.get(row[status_col - 1])
            if fill:
                ws.cell(r, status_col).fill = PatternFill("solid", fgColor=fill)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def main() -> int:
    spec = spec_fields()
    ddl = ddl_columns()
    rows, seen = [], set()
    for table, field, stype, req in spec:
        cols = dict(ddl.get(table, []))
        if field in cols:
            note = NOTES.get((table, field), "")
            status = "Adjusted" if (table, field) in (("bd_leads", "converted_client_id"),
                                                      ("bd_onboarding_handovers", "client_id")) else "As spec"
            rows.append([table, field, stype, req, f"{table}.{field}", short_type(cols[field]), status, note])
        else:
            rows.append([table, field, stype, req, "", "", "Missing", "Not in the draft schema"])
        seen.add((table, field))
    for table, cols in ddl.items():
        for col, defn in cols:
            if (table, col) in seen:
                continue
            reason = ADDED.get((table, col))
            if reason is None and table.startswith("bd_") and table in {t for t, *_ in spec}:
                reason = "Added"
            if reason is None:
                reason = "Supporting table (not one of the 13 spec tables): see Reuse Decisions."
            rows.append([table, "", "", "", f"{table}.{col}", short_type(defn), "Added", reason])
    missing = [r for r in rows if r[6] == "Missing"]

    wb = Workbook()
    lg = wb.active
    lg.title = "Legend"
    lg.column_dimensions["A"].width = 24
    lg.column_dimensions["B"].width = 100
    legend = [
        ("BD Field Mapping", None),
        ("Purpose", "Maps every field of the BD backend specification (docs/BD_BACKEND_SPEC.md) to the draft "
                    "PostgreSQL schema (migrations/bd/001_bd_schema.sql), and BD data to the Client Master at handover."),
        ("Status: As spec", "Column exists with the spec's name and type."),
        ("Status: Adjusted", "Column exists but its type differs from the spec; the Notes column says why."),
        ("Status: Added", "Column not in the spec (supporting tables, generated duplicate-check columns, versioning)."),
        ("Status: Missing", "Spec field with no column. Should be zero."),
        ("Spec fields", str(len(spec))),
        ("Missing", str(len(missing))),
        ("Draft", "Schema not yet applied to any CRM database. See docs/BD_MIGRATION_PLAN.md for the decisions "
                  "needed first."),
        ("Rebuilt by", "Database Module/build_bd_field_mapping.py"),
    ]
    for r, (a, b) in enumerate(legend, start=1):
        ca = lg.cell(r, 1, a)
        ca.font = Font(name=FONT, bold=True, size=16 if r == 1 else 10, color=NAVY if r == 1 else "000000")
        if a.startswith("Status: "):
            ca.fill = PatternFill("solid", fgColor=STATUS_FILL[a[8:]])
        if b is not None:
            cb = lg.cell(r, 2, b)
            cb.font = Font(name=FONT)
            cb.alignment = Alignment(wrap_text=True, vertical="top")
    lg.sheet_view.showGridLines = False

    ws = wb.create_sheet("Spec to Database")
    header(ws, ["Spec Table", "Spec Field", "Spec Type", "Spec Requirement", "Database Column", "Database Type",
                "Status", "Notes"], [26, 28, 18, 38, 40, 44, 11, 60])
    write_rows(ws, rows, {4, 6, 8}, status_col=7)
    dv = DataValidation(type="list", formula1='"As spec,Adjusted,Added,Missing"', allow_blank=True)
    dv.sqref = f"G2:G{ws.max_row}"
    ws.add_data_validation(dv)

    cp = wb.create_sheet("Client Master Prefill")
    header(cp, ["BD Source", "Client Module Column", "Rule"], [58, 26, 80])
    write_rows(cp, [list(r) for r in CLIENT_PREFILL], {1, 3})

    rd = wb.create_sheet("Reuse Decisions")
    header(rd, ["Entity", "Exists in CRM", "Decision", "How"], [24, 52, 14, 70])
    write_rows(rd, [list(r) for r in REUSE], {2, 4})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}: {len(spec)} spec fields, {len(rows) - len(spec)} added columns, {len(missing)} missing")
    for m in missing:
        print("  MISSING", m[0], m[1])
    return 1 if missing else 0


if __name__ == "__main__":
    sys.exit(main())
