"""Plain text out of uploaded files. Nothing here leaves the machine.

PDF text layer (pypdf), Word, Excel, CSV and TXT are read directly. Scanned PDFs (little or no text layer) and
images go through Windows' built-in OCR (ocr.ps1, run with -EncodedCommand so no script file is executed).
"""
from __future__ import annotations

import base64
import csv
import io
import json
import os
import re
import subprocess
import zipfile
from pathlib import Path

PACKAGE_DIR = Path(__file__).resolve().parent

MAX_CHARS = 60_000
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}
OCR_TIMEOUT = 180


def read_text(path: Path, ext: str, max_chars: int = MAX_CHARS) -> tuple[str, str]:
    """(text, method). method: text-layer | ocr | none."""
    ext = ext.lower()
    try:
        if ext == ".pdf":
            text, pages = _pdf(path, max_chars)
            if len(re.sub(r"\s", "", text)) >= 80 * max(1, min(pages, 3)):
                return text, "text-layer"
            ocr = ocr_text(path)
            if len(re.sub(r"\s", "", ocr)) > len(re.sub(r"\s", "", text)):
                return ocr[:max_chars], "ocr"
            return text, "text-layer" if text.strip() else "none"
        if ext in IMAGE_EXT:
            t = ocr_text(path)
            return t[:max_chars], "ocr" if t.strip() else "none"
        if ext in (".txt", ".csv"):
            return _plain(path, max_chars), "text-layer"
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(path, max_chars), "text-layer"
        if ext == ".docx":
            return _docx(path, max_chars), "text-layer"
    except Exception:  # unreadable or malformed files are classified by name only
        return "", "none"
    return "", "none"


def ocr_text(path: Path, max_pages: int = 10) -> str:
    script = (PACKAGE_DIR / "ocr.ps1").read_text(encoding="utf-8")
    encoded = base64.b64encode(script.encode("utf-16-le")).decode("ascii")
    env = {**os.environ, "CM_OCR_PATH": str(path.resolve()), "CM_OCR_MAXPAGES": str(max_pages)}
    try:
        out = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded],
                             capture_output=True, timeout=OCR_TIMEOUT, env=env,
                             creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except (OSError, subprocess.TimeoutExpired):
        return ""
    raw = out.stdout.decode("utf-8", "ignore").strip()
    start = raw.find("[")
    if out.returncode != 0 or start < 0:
        return ""
    try:
        pages = json.loads(raw[start:])
    except json.JSONDecodeError:
        return ""
    return "\n\f\n".join(p for p in pages if isinstance(p, str))


def _pdf(path: Path, max_chars: int) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return "", 0
    out: list[str] = []
    for page in reader.pages[:15]:
        out.append(page.extract_text() or "")
        if sum(map(len, out)) > max_chars:
            break
    return "\n".join(out)[:max_chars], len(reader.pages)


def _plain(path: Path, max_chars: int) -> str:
    raw = path.read_bytes()[: max_chars * 4]
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)[:max_chars]
        except UnicodeDecodeError:
            continue
    return ""


def _xlsx(path: Path, max_chars: int) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out = io.StringIO()
    writer = csv.writer(out)
    try:
        for ws in wb.worksheets[:5]:
            out.write(f"# sheet: {ws.title}\n")
            for row in ws.iter_rows(max_row=400, values_only=True):
                if any(v is not None for v in row):
                    writer.writerow(["" if v is None else v for v in row[:30]])
                if out.tell() > max_chars:
                    break
    finally:
        wb.close()
    return out.getvalue()[:max_chars]


def _docx(path: Path, max_chars: int) -> str:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", "ignore")
    xml = re.sub(r"</w:p>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)[:max_chars]


def zip_member_names(path: Path, limit: int = 2000) -> list[str]:
    try:
        with zipfile.ZipFile(path) as zf:
            return [i.filename for i in zf.infolist()[:limit]]
    except Exception:
        return []
